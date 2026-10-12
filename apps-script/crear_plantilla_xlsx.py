"""Genera la hoja de datos (xlsx -> Google Sheets) del experimento factorial Incentivo x Dificultad.
Uso:  python3 crear_plantilla_xlsx.py [salida.xlsx] [--demo]
  --demo  rellena 40 filas simuladas (solo para comprobar las fórmulas).
Fórmulas solo con SUMPRODUCT, SQRT, FDIST, IFERROR: funcionan igual en Google Sheets y en Excel."""
import json, sys, os, re, random
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

AQUI = os.path.dirname(os.path.abspath(__file__))
H = json.load(open(os.path.join(AQUI, "columnas.json")))
out = next((a for a in sys.argv[1:] if not a.startswith("--")), "Incentivo_Razonamiento_datos.xlsx")
demo = "--demo" in sys.argv
AZUL, MARINO, NARANJA, GRIS = "0047E8", "001447", "FF9606", "43506B"
FN = "Arial"
wb = Workbook()

cfg = open(os.path.join(AQUI, "..", "js", "config.js"), encoding="utf-8").read()
ITEMS = re.findall(r'\{ id: "(\w+)", dificultad: "(\w+)", reglas: \[([^\]]*)\], correcta: (\d+) \}', cfg)
ITEMS = [(i, d, [r.strip().strip('"') for r in rg.split(",")], int(c)) for i, d, rg, c in ITEMS]
EXP = re.findall(r'"(\w+)"', re.search(r"ITEMS_EXPERIMENTO: \[(.*?)\]", cfg, re.S).group(1))
UMBRAL = int(re.search(r"UMBRAL: (\d+)", cfg).group(1))
REGLAS = re.findall(r'(\w+): "([^"]+)"', re.search(r"REGLAS: \{(.*?)\}", cfg, re.S).group(1))


def cab(ws, fila, textos, color=AZUL, col0=1):
    for i, t in enumerate(textos, col0):
        c = ws.cell(fila, i, t)
        c.font = Font(name=FN, bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=color)
        c.alignment = Alignment(vertical="center", wrap_text=True)


def nota(ws, celda, texto, rango=None, alto=None):
    ws[celda] = texto; ws[celda].font = Font(name=FN, size=9, italic=True, color=GRIS)
    ws[celda].alignment = Alignment(wrap_text=True, vertical="top")
    if rango: ws.merge_cells(rango)
    if alto: ws.row_dimensions[ws[celda].row].height = alto


# ---------------- Respuestas ----------------
ws = wb.active; ws.title = "Respuestas"
cab(ws, 1, H); ws.freeze_panes = "B2"
for i, h in enumerate(H, 1):
    ws.column_dimensions[get_column_letter(i)].width = max(9, min(22, len(h) + 3))
col = {h: get_column_letter(i + 1) for i, h in enumerate(H)}
N = 3000
def R(h): return f"Respuestas!${col[h]}${2}:${col[h]}${N}"

if demo:
    random.seed(7)
    for r in range(2, 42):
        cond = "incentivo" if r % 2 else "control"
        hab = random.gauss(0, 1)
        fila = {h: "" for h in H}
        fila.update(id=f"DEMO-{r - 1}", cohorte="MDTI-2026-G1", modo="experimento", condicion=cond, reclutador=f"R{r % 5}",
                    sospechoso=1 if r in (7, 20) else 0, ya_participo="Sí" if r == 11 else "No",
                    condiciones="No" if r == 15 else "Sí", recuerda_premio="Sí" if cond == "incentivo" else "No")
        f = max(0, min(6, round(4 + hab + (0.4 if cond == "incentivo" else 0) + random.gauss(0, .8))))
        d = max(0, min(6, round(2 + hab + (0.9 if cond == "incentivo" else 0) + random.gauss(0, .8))))
        fila.update(aciertos_facil=f, aciertos_dificil=d, aciertos_total=f + d, n_facil=6, n_dificil=6,
                    tr_mediana_facil=random.randint(6000, 14000), tr_mediana_dificil=random.randint(12000, 26000),
                    supera_umbral=int(f + d >= UMBRAL), premio=int(cond == "incentivo" and f + d >= UMBRAL))
        for k in ("mot_esfuerzo", "mot_interes", "mot_importancia"):
            fila[k] = max(1, min(5, round(3.3 + (0.6 if cond == "incentivo" else 0) + random.gauss(0, .9))))
        fila["motivacion"] = round((fila["mot_esfuerzo"] + fila["mot_interes"] + fila["mot_importancia"]) / 3, 3)
        for it, dif, rg, cor in ITEMS:
            if it in EXP:
                ok = int(random.random() < (0.75 if dif == "facil" else 0.4)); fila[it + "_ok"] = ok
                fila[it + "_resp"] = cor if ok else random.randint(1, 8); fila[it + "_tr"] = random.randint(3000, 30000)
        for i, h in enumerate(H, 1): ws.cell(r, i, fila[h])

# ---------------- Resumen (fórmulas en vivo) ----------------
rs = wb.create_sheet("Resumen", 0)
rs.column_dimensions["A"].width = 46
for c in "BCDEFGH": rs.column_dimensions[c].width = 15
rs["A1"] = "Incentivo × Dificultad en razonamiento · Resumen en vivo"
rs["A1"].font = Font(name=FN, size=16, bold=True, color=AZUL)
rs["A2"] = "ESIC University · Métodos, Diseño y Técnicas de Investigación en Psicología · Tema 6 (diseños factoriales)"
rs["A2"].font = Font(name=FN, size=10, color=MARINO)
rs["A4"] = "Cohorte a analizar (vacío = todas):"; rs["B4"] = ""
rs["B4"].fill = PatternFill("solid", fgColor="FFF3E0"); rs["B4"].font = Font(name=FN, bold=True)
nota(rs, "A5", "Se excluyen automáticamente: participaciones marcadas como sospechosas (sospechoso = 1), quienes ya habían hecho la tarea "
     "y quienes declaran interrupciones. El análisis completo (efectos simples, tiempos, matriz Q, calidad de datos) está en la pestaña "
     "«Resultados» de la web.", "A5:H5", 32)

FA, DI = R("aciertos_facil"), R("aciertos_dificil")
MASK = (f'({FA}<>"")*({DI}<>"")*({R("sospechoso")}<>1)*({R("ya_participo")}<>"Sí")*({R("condiciones")}<>"No")'
        f'*((({R("cohorte")}=$B$4)+($B$4=""))>0)')
G = {"incentivo": f'{MASK}*({R("condicion")}="incentivo")', "control": f'{MASK}*({R("condicion")}="control")'}

cab(rs, 7, ["Descriptivos (aciertos)", "n", "Fácil: media", "Fácil: DE", "Difícil: media", "Difícil: DE", "Total: media", "≥ umbral (%)"])
filas = {"incentivo": 8, "control": 9}
for g, r in filas.items():
    m = G[g]; rs[f"A{r}"] = "Incentivo" if g == "incentivo" else "Control"
    rs[f"B{r}"] = f"=SUMPRODUCT({m})"
    rs[f"C{r}"] = f'=IFERROR(SUMPRODUCT({m}*{FA})/B{r},"")'
    rs[f"D{r}"] = f'=IFERROR(SQRT(SUMPRODUCT({m}*({FA}-C{r})^2)/(B{r}-1)),"")'
    rs[f"E{r}"] = f'=IFERROR(SUMPRODUCT({m}*{DI})/B{r},"")'
    rs[f"F{r}"] = f'=IFERROR(SQRT(SUMPRODUCT({m}*({DI}-E{r})^2)/(B{r}-1)),"")'
    rs[f"G{r}"] = f'=IFERROR(C{r}+E{r},"")'
    rs[f"H{r}"] = f'=IFERROR(SUMPRODUCT({m}*(({FA}+{DI})>=$B$12))/B{r},"")'
    rs[f"H{r}"].number_format = "0%"
    for c in "CDEFG": rs[f"{c}{r}"].number_format = "0.00"
rs["A10"] = "Total"; rs["B10"] = "=B8+B9"
rs["C10"] = '=IFERROR((B8*C8+B9*C9)/B10,"")'; rs["E10"] = '=IFERROR((B8*E8+B9*E9)/B10,"")'; rs["G10"] = '=IFERROR(C10+E10,"")'
for c in "CEG": rs[f"{c}10"].number_format = "0.00"
rs["A12"] = "Umbral del premio (aciertos)"; rs["B12"] = UMBRAL
rs["A13"] = "Cafés a pagar (grupo incentivo que supera el umbral)"; rs["B13"] = f'=SUMPRODUCT({G["incentivo"]}*(({FA}+{DI})>=$B$12))'

# Cálculos auxiliares del ANOVA mixto: s = (fácil + difícil)/2 ; d = fácil - difícil
cab(rs, 15, ["Auxiliares por grupo", "n", "media s", "SC dentro s", "media d", "SC dentro d"], MARINO)
for g, r in (("incentivo", 16), ("control", 17)):
    m = G[g]; rs[f"A{r}"] = "Incentivo" if g == "incentivo" else "Control"; rs[f"B{r}"] = f"=B{filas[g]}"
    rs[f"C{r}"] = f'=IFERROR(SUMPRODUCT({m}*({FA}+{DI})/2)/B{r},"")'
    rs[f"D{r}"] = f'=IFERROR(SUMPRODUCT({m}*(({FA}+{DI})/2-C{r})^2),"")'
    rs[f"E{r}"] = f'=IFERROR(SUMPRODUCT({m}*({FA}-{DI}))/B{r},"")'
    rs[f"F{r}"] = f'=IFERROR(SUMPRODUCT({m}*(({FA}-{DI})-E{r})^2),"")'
    for c in "CDEF": rs[f"{c}{r}"].number_format = "0.000"
rs["A18"] = "Medias globales (ponderadas)"; rs["C18"] = '=IFERROR((B16*C16+B17*C17)/(B16+B17),"")'; rs["E18"] = '=IFERROR((B16*E16+B17*E17)/(B16+B17),"")'
for c in "CE": rs[f"{c}18"].number_format = "0.000"

cab(rs, 20, ["ANOVA mixto 2 × 2 (SC tipo III, como JASP)", "SC", "gl", "MC", "F", "p", "η² parcial"])
rs["A21"] = "Dificultad (intra)"
rs["A22"] = "Dificultad × Incentivo"
rs["A23"] = "Residuo intra"
rs["A24"] = "Incentivo (entre)"
rs["A25"] = "Residuo entre"
glr = "(B16+B17-2)"
rs["C23"] = f'=IFERROR({glr},"")'; rs["B23"] = '=IFERROR((F16+F17)/2,"")'; rs["D23"] = '=IFERROR(B23/C23,"")'
# F de la dificultad: media NO ponderada de d frente a su error típico
rs["E21"] = f'=IFERROR(((E16+E17)/2)^2/(((F16+F17)/{glr})*(1/B16+1/B17)/4),"")'
rs["C21"] = 1; rs["B21"] = '=IFERROR(E21*D23,"")'; rs["D21"] = '=IFERROR(B21/C21,"")'
rs["B22"] = '=IFERROR((B16*(E16-E18)^2+B17*(E17-E18)^2)/2,"")'; rs["C22"] = 1; rs["D22"] = '=IFERROR(B22/C22,"")'; rs["E22"] = '=IFERROR(D22/D23,"")'
rs["B24"] = '=IFERROR(2*(B16*(C16-C18)^2+B17*(C17-C18)^2),"")'; rs["C24"] = 1; rs["D24"] = '=IFERROR(B24/C24,"")'
rs["B25"] = '=IFERROR(2*(D16+D17),"")'; rs["C25"] = f'=IFERROR({glr},"")'; rs["D25"] = '=IFERROR(B25/C25,"")'; rs["E24"] = '=IFERROR(D24/D25,"")'
for r in (21, 22, 24):
    err = 25 if r == 24 else 23
    rs[f"F{r}"] = f'=IFERROR(FDIST(E{r},C{r},C{err}),"")'
    rs[f"G{r}"] = f'=IFERROR(B{r}/(B{r}+B{err}),"")'
for r in range(21, 26):
    for c in "BDEFG": rs[f"{c}{r}"].number_format = "0.000"

cab(rs, 27, ["Comprobación de la manipulación", "Incentivo", "Control"])
for k, (campo, nom) in enumerate([("mot_esfuerzo", "Esfuerzo (1-5)"), ("mot_interes", "Interés (1-5)"),
                                   ("mot_importancia", "Importancia (1-5)"), ("motivacion", "Motivación (media)")]):
    r = 28 + k; rs[f"A{r}"] = nom
    for c, g in (("B", "incentivo"), ("C", "control")):
        rs[f"{c}{r}"] = f'=IFERROR(SUMPRODUCT({G[g]}*({R(campo)}<>""),{R(campo)})/SUMPRODUCT({G[g]}*({R(campo)}<>"")),"")'
        rs[f"{c}{r}"].number_format = "0.00"
rs["A32"] = "Dice que se le ofreció premio (%)"
for c, g in (("B", "incentivo"), ("C", "control")):
    rs[f"{c}32"] = f'=IFERROR(SUMPRODUCT({G[g]}*({R("recuerda_premio")}="Sí"))/SUMPRODUCT({G[g]}),"")'; rs[f"{c}32"].number_format = "0%"
nota(rs, "A34", "Las fórmulas toleran grupos de distinto tamaño (SC de tipo III). Los resultados deben coincidir con los de la web y con JASP "
     "(ANOVA de medidas repetidas: factor intra «Dificultad» con 2 niveles; factor entre «Incentivo»).", "A34:H34", 30)
for row in rs.iter_rows(min_row=8, max_row=32):
    for c in row:
        if c.fill.fgColor.rgb not in ("00" + AZUL, "00" + MARINO):
            c.font = Font(name=FN, bold=c.font.bold)

# ---------------- Asignaciones ----------------
asg = wb.create_sheet("Asignaciones")
cab(asg, 1, ["fecha", "cohorte", "id", "condicion", "bloque", "quedan_en_bloque"])
for c, w in zip("ABCDEF", [20, 16, 20, 12, 9, 16]): asg.column_dimensions[c].width = w

# ---------------- Items: matriz Q + p en vivo ----------------
it = wb.create_sheet("Items")
cab(it, 1, ["Ítem", "Dificultad", "En el experimento"] + [n for _, n in REGLAS] + ["Nº reglas", "Clave", "N", "p (aciertos)", "p incentivo", "p control"])
it.column_dimensions["A"].width = 8; it.column_dimensions["B"].width = 11; it.column_dimensions["C"].width = 12
for j in range(4, 4 + len(REGLAS)): it.column_dimensions[get_column_letter(j)].width = 15
it.row_dimensions[1].height = 44
for i, (iid, dif, rg, cor) in enumerate(ITEMS, 2):
    vals = [iid, {"practica": "Práctica", "facil": "Fácil", "dificil": "Difícil"}[dif], "Sí" if iid in EXP else ("Práctica" if dif == "practica" else "No")]
    vals += [1 if k in rg else 0 for k, _ in REGLAS] + [len(rg), cor]
    for j, v in enumerate(vals, 1):
        c = it.cell(i, j, v); c.font = Font(name=FN)
        if 4 <= j < 4 + len(REGLAS):
            c.alignment = Alignment(horizontal="center")
            if v == 1: c.fill = PatternFill("solid", fgColor="EAF0FF"); c.font = Font(name=FN, bold=True, color=AZUL)
    j0 = 6 + len(REGLAS)
    if dif != "practica" and iid + "_ok" in col:
        ok = R(iid + "_ok")
        it.cell(i, j0, f'=SUMPRODUCT({MASK}*({ok}<>""))')
        L = get_column_letter(j0)
        it.cell(i, j0 + 1, f'=IFERROR(SUMPRODUCT({MASK}*({ok}<>""),{ok})/{L}{i},"")').number_format = "0.00"
        for k, g in enumerate(("incentivo", "control")):
            it.cell(i, j0 + 2 + k, f'=IFERROR(SUMPRODUCT({G[g]}*({ok}<>""),{ok})/SUMPRODUCT({G[g]}*({ok}<>"")),"")').number_format = "0.00"
it.freeze_panes = "B2"
nota(it, f"A{len(ITEMS) + 3}", "Matriz Q: 1 = el ítem usa esa regla. Fáciles: 1 regla. Difíciles: rotación de la figura principal + sustracción de líneas + una tercera regla. "
     "Ítems generados con IMak (Blum y Holling, 2018).", f"A{len(ITEMS) + 3}:{get_column_letter(j0 + 3)}{len(ITEMS) + 3}", 30)

# ---------------- Léeme ----------------
lm = wb.create_sheet("Leeme")
lm.column_dimensions["A"].width = 110
txt = [
    ("Cómo funciona esta hoja", True),
    ("• «Respuestas»: una fila por participación. La escribe el Apps Script de la web; no edites ni reordenes las columnas.", False),
    ("• «Asignaciones»: registro de la aleatorización (bloques permutados de 4 por cohorte). Sirve para auditar que nadie elige su condición.", False),
    ("• «Resumen»: descriptivos, ANOVA mixto 2 × 2 y comprobación de la manipulación, en vivo. Escribe una cohorte en B4 para filtrar.", False),
    ("• «Items»: matriz Q (reglas de cada ítem) y proporción de aciertos de cada ítem.", False),
    ("", False),
    ("Concurrencia: si dos personas envían a la vez, el script usa LockService y escribe una fila detrás de otra; además ignora filas con un id repetido.", False),
    ("Tiempos: <ítem>_tr = milisegundos desde que aparecen las opciones hasta la respuesta; <ítem>_resp: 1-8 = opción, 9 = ninguna, 10 = no sé, 0 = tiempo agotado.", False),
    ("Sospechoso = 1 si: muchas respuestas < 2,5 s, mediana muy baja, misma opción repetida, duración muy corta o salidas de la pestaña (motivo en motivo_sospecha).", False),
    ("Privacidad: no se recogen nombres ni correos. Los códigos de premio (CAFE-xxxx) solo sirven para canjear el café.", False),
]
for i, (t, b) in enumerate(txt, 1):
    c = lm.cell(i, 1, t); c.font = Font(name=FN, bold=b, size=13 if b else 10, color=AZUL if b else "000000"); c.alignment = Alignment(wrap_text=True)

wb.save(out)
print("guardado", out, "· columnas:", len(H), "· ítems:", len(ITEMS))
